"""
ingest_gate1_answers.py
═══════════════════════════════════════════════════════════════════════════

把闸一标注表（build_gate1_human_sheets.py 生成、填完发回的 xlsx）落进
truth_vault.note_feature_answers —— 人填的和别的模型填的都走这一条（A15, D-079 §5.1）。

    python ingest_gate1_answers.py --extractor human:张三 --file 笔记标注_标注员A.xlsx
    python ingest_gate1_answers.py --extractor jev:1.13.0-A --file A_Jev.xlsx   # 一张表一个 extractor, 见下
    python ingest_gate1_answers.py ... --write        # 默认只校验不写

extractor 的写法（与 annotate_feature_pass 的 llm:<model> / code:v1 并列）:
    human:<姓名>     一个人一张表, 姓名进 extractor 才分得清谁和谁不一致（κ 按人算）
    jev:<版本>-<表>  别的模型填的表 —— 它【不是】闸一"人工那一半", 只是又一个模型。
                     一个模型填了 A/B/C 三张, 重叠的 50 篇它答了两遍, 主键里没有"哪张表",
                     不带表的字母就会互相覆盖, 所以 extractor 要写成 jev:1.13.0-A 并一张一跑
run_tag 默认 gate1-20260928, 与 D-079 一致; 主键 (subject, question, version, extractor, run_tag)
决定同一个人同一张表重跑是覆盖, 不是再插一份。

校验（任一条不过就整份不写, 别把半张表落库）:
    · 表里的 note_id 必须全在名单 CSV 里, 且恰好等于名单里分给这位标注员的那 50 篇
      （名单 `annotators` 列）—— 发错表、混了表、少行多行都在这里拦住
    · 表头的中文短题名必须能一一映射回题库 question_id（题库改了题名这里会红）
    · 答案必须在闭集里（bool → 是/否; choice → options）；「—」是预先灰掉的格子, 跳过
      且必须与名单 `skipped_questions` 一致；空格子按【没答】计, 不写行, 但会点名
      名单那一格用 gate1_labels.split_skipped 拆 ("|" 与旧的 "," 都认), 拆出题库里没有的题号直接红
      (D-085: 以前按 "," 拆 "|" 拼的名单, 11 个灰格拆成一个认不出的长串, 两个方向的校验都失效)
    · 「拿不准的地方」那一列不进库; 只有 Jev 那种按题写了概率的, 解析出来写进 prob

prob 的口径 (D-085, 对齐 notes_v1_17): **所选答案的概率**。Jev 在是非题上写的是「判「是」的概率 p」,
答「是」存 p、答「否」存 1 − p; 选择题写的是「第一名只有 p」, Jev 的答案就是第一名, 存 p。
D-081 那三张表 (run_tag gate1-20260928) 是按旧口径落的 (是非题存 P(是)), 不回改; D-084 已定 Jev 不重跑,
真要重收请落新的 run_tag, 别把两种口径混进同一个 run_tag。
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path

import feature_bank as fb
from _common import _iso_now, get_supabase_client, setup_logger
from gate1_labels import SHORT, split_skipped

logger = setup_logger("ingest_gate1")

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_MANIFEST = ROOT / "data-analysis" / "gate1-human-sample-2026-09-28.csv"
DEFAULT_RUN_TAG = "gate1-20260928"
SHEET = "标注"
SKIP_MARK = "—"
UPSERT_CHUNK = 200                       # 走 body 不走 URL, 但一次别太大

_EXTRACTOR_RE = re.compile(r"^(human|jev):[^\s:]+$")
# Jev 在「拿不准的地方」里的写法: "4 具体地点或场合：判「是」的概率 0.39，落在 0.35–0.65"
#                                "12 产品角色：第一名只有 0.45；「主角」0.45 与「只暗示」0.33 太接近"
# 两种写法的数不是一个东西: 前者是 P(是), 后者是第一名 (= Jev 选的那个) 的概率。分开记, build_rows 换算。
_PROB_RE = re.compile(r"(?m)^(\d+) [^：]+：(判「是」的概率|第一名只有) ([0-9.]+)")
PROB_YES, PROB_TOP = "p_yes", "p_top"


def parse_probs(text: str, num_to_qid: dict[int, str]) -> dict[str, tuple[str, float]]:
    """「拿不准的地方」那一格 → {qid: (PROB_YES | PROB_TOP, 数)}; 题号对不上的行、不在 [0,1] 的数忽略。"""
    out: dict[str, tuple[str, float]] = {}
    for n, kind, p in _PROB_RE.findall(str(text or "")):
        try:
            v = float(p)
        except ValueError:
            continue
        if int(n) in num_to_qid and 0.0 <= v <= 1.0:
            out[num_to_qid[int(n)]] = (PROB_YES if kind.startswith("判") else PROB_TOP, v)
    return out


def prob_of_answer(answer: str, prob) -> float | None:
    """(PROB_YES | PROB_TOP, p) + 表里的答案 → 所选答案的概率 (v1.17 口径)。

    PROB_YES 只对是非题有意义: 答「是」→ p, 答「否」→ 1 − p (四舍五入到 4 位, 免得 0.61000000001)。
    PROB_TOP 是 Jev 第一名的概率, Jev 的答案就是第一名 → p。
    PROB_YES 却配了非是非答案 (对不上) → None, 宁可不写也不写错。"""
    if prob is None:
        return None
    kind, p = prob
    if kind == PROB_TOP:
        return p
    if answer == "是":
        return p
    if answer == "否":
        return round(1.0 - p, 4)
    return None


def load_manifest(path: Path) -> dict[str, dict]:
    with open(path, encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    if not rows:
        raise SystemExit(f"名单为空: {path}")
    return {r["note_id"]: r for r in rows}


def read_sheet(path: Path, bank: dict) -> tuple[str, list[dict]]:
    """返回 (标注员字母, [{note_id, answers: {qid: 值}, probs: {qid: (PROB_YES|PROB_TOP, float)}}])。"""
    import openpyxl   # 只有读 xlsx 要它; 放这里, 校验 / 拼行这些纯函数不装它也能 import (CI 就不装)
    wb = openpyxl.load_workbook(path, data_only=True)
    if SHEET not in wb.sheetnames:
        raise SystemExit(f"{path.name}: 没有「{SHEET}」页, 页有 {wb.sheetnames}")
    # 标注员字母在「怎么标」页的标题里: 笔记标注 · 标注员 A · 共 50 篇
    who = None
    if "怎么标" in wb.sheetnames:
        for row in wb["怎么标"].iter_rows(values_only=True):
            for v in row:
                m = re.search(r"标注员 ([ABC])", str(v or ""))
                if m:
                    who = m.group(1)
                    break
            if who:
                break
    if not who:
        raise SystemExit(f"{path.name}: 「怎么标」页里找不到「标注员 A/B/C」, 分不清这是谁的表")

    ws = wb[SHEET]
    header = [c.value for c in ws[1]]
    label_to_qid = {v: k for k, v in SHORT.items()}
    qidx = fb.question_index(bank)
    col_q: dict[int, str] = {}
    unc_col = nid_col = None
    for i, h in enumerate(header):
        if h is None:
            continue
        h = str(h)
        if h == "note_id":
            nid_col = i
            continue
        if h.startswith("拿不准"):
            unc_col = i
            continue
        m = re.match(r"^\d+\n(.+)$", h)
        if not m:
            continue
        label = m.group(1).strip()
        if label not in label_to_qid:
            raise SystemExit(f"{path.name}: 表头「{label}」映射不回题库 question_id (SHORT 里没有)")
        qid = label_to_qid[label]
        if qid not in qidx:
            raise SystemExit(f"{path.name}: 题库里没有 {qid}")
        col_q[i] = qid
    if nid_col is None:
        raise SystemExit(f"{path.name}: 没有隐藏的 note_id 列, 回收对不了账")
    if len(col_q) != len(fb.llm_questions(bank)):
        raise SystemExit(f"{path.name}: 映射到 {len(col_q)} 题, 题库有 {len(fb.llm_questions(bank))} 题")

    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        nid = row[nid_col]
        if not nid:
            continue
        answers = {qid: (None if row[i] is None else str(row[i]).strip()) for i, qid in col_q.items()}
        probs: dict[str, tuple[str, float]] = {}
        if unc_col is not None and row[unc_col]:
            num_to_qid = {int(h.split("\n")[0]): col_q[i] for i, h in enumerate(header) if i in col_q}
            probs = parse_probs(str(row[unc_col]), num_to_qid)
        out.append({"note_id": str(nid), "answers": answers, "probs": probs})
    return who, out


def validate(path: Path, who: str, notes: list[dict], manifest: dict, bank: dict) -> list[str]:
    errs: list[str] = []
    # 名单里 annotators 写成 "AB" / "AC" / "B" / "C" (字母串, 不带分隔符)
    expected = {nid for nid, r in manifest.items() if who in (r["annotators"] or "")}
    got = [n["note_id"] for n in notes]
    if len(got) != len(set(got)):
        errs.append("note_id 有重复行")
    if set(got) != expected:
        errs.append(f"标注员 {who} 的表应有名单里的 {len(expected)} 篇, 实得 {len(got)} 篇; "
                    f"多出 {sorted(set(got) - expected)[:3]}, 缺 {sorted(expected - set(got))[:3]}")
    qidx = fb.question_index(bank)
    for n in notes:
        man = manifest.get(n["note_id"])
        if not man:
            continue
        skipped = set(split_skipped(man.get("skipped_questions")))
        # 拆出题库里没有的题号 = 名单坏了或分隔符又对不上了 (D-085 那 11 格就是这么漏进库的)。
        # 这时灰格校验在两个方向上都是空转, 必须直接红, 不能当「这篇没有灰格」。
        unknown = sorted(skipped - set(qidx))
        if unknown:
            errs.append(f"{n['note_id']}: 名单 skipped_questions 里有题库没有的题号 {unknown} —— 名单坏了或分隔符不对")
            continue
        for qid, v in n["answers"].items():
            if v == SKIP_MARK:
                if qid not in skipped:
                    errs.append(f"{n['note_id']} {qid}: 表里是「—」但名单没说这题不问")
                continue
            if qid in skipped:
                errs.append(f"{n['note_id']} {qid}: 名单说这题不问(灰格), 表里却有答案 {v!r}")
                continue
            if v in (None, ""):
                continue                       # 没答, 汇总时点名
            if v not in fb.closed_set(qidx[qid]):
                errs.append(f"{n['note_id']} {qid}: 答案 {v!r} 不在闭集 {fb.closed_set(qidx[qid])}")
    return errs


def build_rows(notes: list[dict], bank: dict, extractor: str, run_tag: str) -> tuple[list[dict], list[str]]:
    qidx = fb.question_index(bank)
    rows, blanks = [], []
    now = _iso_now()
    for n in notes:
        for qid, v in n["answers"].items():
            if v == SKIP_MARK:
                continue
            if v in (None, ""):
                blanks.append(f"{n['note_id']}:{qid}")
                continue
            rows.append({
                "subject_type": "note", "subject_id": n["note_id"],
                "question_id": qid, "question_version": int(qidx[qid].get("version") or 1),
                "bank_version": bank["bank_version"], "bank_sha256": bank["_sha256"],
                "extractor": extractor, "run_tag": run_tag,
                "answer": v, "evidence": None, "prob": prob_of_answer(v, n["probs"].get(qid)),
                "invalid_reason": None, "extracted_at": now,
            })
    return rows, blanks


_SQL_UPSERT_TAIL = (
    "on conflict (subject_type, subject_id, question_id, question_version, extractor, run_tag) "
    "do update set answer = excluded.answer, evidence = excluded.evidence, prob = excluded.prob, "
    "invalid_reason = excluded.invalid_reason, bank_version = excluded.bank_version, "
    "bank_sha256 = excluded.bank_sha256, extracted_at = excluded.extracted_at;")


def _sql_lit(v) -> str:
    if v is None:
        return "NULL"
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return repr(v)
    return "'" + str(v).replace("'", "''") + "'"


def rows_to_sql(rows: list[dict]) -> str:
    """与 --write 的 upsert 等价的 SQL: 主键冲突就覆盖 (同一张表重跑 = 覆盖, 不是再插)。

    形态是【一篇一行, 20 题两个数组】再 unnest 展开, 不是一格一行: 一张表 1,000 格摊平要
    230 KB, 这样 ~10 KB —— 拿到 MCP / SQL 编辑器里才贴得动。共同的列 (extractor / run_tag /
    题库版本 / 时间) 只写一次, 与 --write 的每行完全等价。
    """
    if not rows:
        return "-- ingest_gate1_answers.py · nothing to write\n"
    # 共同列必须真的共同 —— 一份 SQL 对应一次调用 (一张表, 一个 extractor)
    common = {k: {r[k] for r in rows} for k in ("extractor", "run_tag", "bank_version", "bank_sha256",
                                                "extracted_at", "subject_type")}
    bad = {k: v for k, v in common.items() if len(v) != 1}
    if bad:
        raise SystemExit(f"rows_to_sql: 这些列在一份 SQL 里不该有多个值: {bad}")
    c = {k: next(iter(v)) for k, v in common.items()}
    qids = list(dict.fromkeys(r["question_id"] for r in rows))          # 保持题序
    qver = {r["question_id"]: r["question_version"] for r in rows}
    per_subject: dict[str, dict[str, dict]] = {}
    for r in rows:
        per_subject.setdefault(r["subject_id"], {})[r["question_id"]] = r
    values = []
    for sid, byq in per_subject.items():
        ans = ", ".join(_sql_lit(byq[q]["answer"]) if q in byq else "NULL" for q in qids)
        prob = ", ".join(_sql_lit(byq[q]["prob"]) if q in byq else "NULL" for q in qids)
        values.append(f"  ({_sql_lit(sid)}, array[{ans}], array[{prob}]::real[])")
    q_ids = ", ".join(_sql_lit(q) for q in qids)
    q_ver = ", ".join(str(int(qver[q])) for q in qids)
    return (
        f"-- ingest_gate1_answers.py · note_feature_answers upsert · {len(rows)} 格 · "
        f"{len(per_subject)} 篇 · extractor={c['extractor']} run_tag={c['run_tag']} · generated {_iso_now()}\n"
        f"with q as (\n  select * from unnest(array[{q_ids}], array[{q_ver}]) with ordinality "
        "as t(question_id, question_version, ord)\n),\nd(subject_id, answers, probs) as (values\n"
        + ",\n".join(values) + "\n)\n"
        "insert into truth_vault.note_feature_answers (subject_type, subject_id, question_id, question_version, "
        "bank_version, bank_sha256, extractor, run_tag, answer, evidence, prob, invalid_reason, extracted_at)\n"
        f"select {_sql_lit(c['subject_type'])}, d.subject_id, q.question_id, q.question_version, "
        f"{_sql_lit(c['bank_version'])}, {_sql_lit(c['bank_sha256'])}, {_sql_lit(c['extractor'])}, "
        f"{_sql_lit(c['run_tag'])}, d.answers[q.ord], NULL, d.probs[q.ord], NULL, {_sql_lit(c['extracted_at'])}\n"
        "from d cross join q\nwhere d.answers[q.ord] is not null   -- 灰格 (题库说这题不问) 没有行\n"
        + _SQL_UPSERT_TAIL + "\n")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[1])
    ap.add_argument("--file", action="append", required=True, type=Path, help="填完的 xlsx, 可重复")
    ap.add_argument("--extractor", required=True, help="human:<姓名> 或 jev:<版本>")
    ap.add_argument("--run-tag", default=DEFAULT_RUN_TAG)
    ap.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    ap.add_argument("--write", action="store_true", help="真的写库; 不给就只校验")
    ap.add_argument("--sql-out", type=Path, default=None,
                    help="不直接写库, 把同一批行写成 INSERT … ON CONFLICT DO UPDATE 的 .sql "
                         "(没有 service key 的机器上拿去 Supabase SQL 编辑器 / MCP 执行)")
    args = ap.parse_args()

    if not _EXTRACTOR_RE.match(args.extractor):
        raise SystemExit(f"extractor 要写成 human:<姓名> 或 jev:<版本>, 实得 {args.extractor!r}")
    if args.extractor.startswith("human:") and len(args.file) > 1:
        raise SystemExit("human:<姓名> 一个人只有一张表; 多张表请分开跑, 各写各的名字")

    bank = fb.load_bank()
    manifest = load_manifest(args.manifest)
    all_rows: list[dict] = []
    for path in args.file:
        who, notes = read_sheet(path, bank)
        errs = validate(path, who, notes, manifest, bank)
        if errs:
            for e in errs[:20]:
                logger.error("%s: %s", path.name, e)
            raise SystemExit(f"{path.name}: {len(errs)} 处校验不过, 整份不写")
        rows, blanks = build_rows(notes, bank, args.extractor, args.run_tag)
        with_prob = sum(1 for r in rows if r["prob"] is not None)
        logger.info("%s: 标注员 %s · %d 篇 · %d 格答案 · %d 格没答 · %d 格带概率",
                    path.name, who, len(notes), len(rows), len(blanks), with_prob)
        if blanks:
            logger.warning("%s: 没答的格子(不写行): %s%s", path.name, ", ".join(blanks[:10]),
                           " …" if len(blanks) > 10 else "")
        all_rows.extend(rows)

    # 主键里没有"哪张表": 同一 extractor 下同一篇答了两次(A/B、A/C 各重叠 25 篇), 后写的会把
    # 先写的盖掉, 双标那一半就丢了。人是一人一名字不会撞; 一个模型填三张表就会 —— 要它保留
    # 重叠, extractor 得带上表的字母(jev:1.13.0-A / -B / -C), 每张单独跑。
    seen: dict[tuple[str, str], int] = {}
    for r in all_rows:
        seen[(r["subject_id"], r["question_id"])] = seen.get((r["subject_id"], r["question_id"]), 0) + 1
    dup = sorted({k[0] for k, c in seen.items() if c > 1})
    if dup:
        raise SystemExit(f"{len(dup)} 篇在多张表里都答了(如 {dup[:3]}), 同一 extractor={args.extractor!r} 下会互相覆盖; "
                         "给每张表不同的 extractor(如 jev:1.13.0-A)并分开跑")

    logger.info("合计 %d 行 → note_feature_answers (extractor=%s, run_tag=%s)%s",
                len(all_rows), args.extractor, args.run_tag, "" if args.write else " [dry-run, 没写]")
    if args.sql_out:
        sql = rows_to_sql(all_rows)
        args.sql_out.write_text(sql, encoding="utf-8")
        logger.info("SQL 写到 %s (%d 格, %d 字节)", args.sql_out, len(all_rows), len(sql.encode("utf-8")))
    if not args.write:
        return 0
    sb = get_supabase_client()
    for i in range(0, len(all_rows), UPSERT_CHUNK):
        sb.schema("truth_vault").table("note_feature_answers").upsert(all_rows[i:i + UPSERT_CHUNK]).execute()
    logger.info("写完 %d 行", len(all_rows))
    return 0


if __name__ == "__main__":
    sys.exit(main())
