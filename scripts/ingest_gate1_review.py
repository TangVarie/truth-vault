#!/usr/bin/env python3
"""ingest_gate1_review.py —— owner 填回的闸一裁决表 → note_feature_answers 的 SQL (D-128)。

    cd scripts && python ingest_gate1_review.py --file /path/to/scratchpad/闸一裁决.xlsx \
        --cells ../data-analysis/gate1-reader-cells-<日期>.csv --sql-out /path/to/scratchpad/owner.sql

只出 SQL, 不写库: 这里的 MCP 是只读的, SQL 交给有写权限的人执行 (ingest_gate1_answers.rows_to_sql 的形状:
一篇一行、两题两个数组再 unnest, 主键冲突就覆盖 —— 重收同一张表是覆盖, 不是再插)。

写成什么:
  · extractor = human:owner, run_tag = fqr-g1-owner (fb.check_run_tag: primary / gate1-* / 不带 fqr- 都拒;
    也不许和模型的三趟重名), question_version / bank_version / bank_sha256 = 副题库的;
  · answer = owner 在「你的判断」里选的值; evidence / prob = NULL —— 「理由」列不进库 (可能抄了原文, D-110 的先例)。
校验 (任一条不过整份不出, 别把半张表落库):
  · 每个 cell_id 都在格子表里; 题页上的格 kind=disagree 且题号就是这一页的题; 「抽查」页的格 kind=spot;
  · 同一个 cell_id 不许出现两次; 答案必须在那道题的闭集里; 「拿不准」和空着的不写行, 只计数。
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

import feature_bank as fb
from _common import _iso_now
from build_gate1_review_sheet import SPOT_SHEET, UNSURE, sheet_name
from gate1_labels import SHORT
from gate1_models import DEFAULT_RUNS, OWNER_EXTRACTOR, OWNER_RUN_TAG, Gate1Refused, load_secondary, read_cells
from ingest_gate1_answers import rows_to_sql

PICK_HEADER = "你的判断"


class IngestRefused(Exception):
    pass


def read_review(path: Path, bank: dict) -> list[dict]:
    """xlsx → [{sheet, cell_id, value}] (题页 + 抽查页; 怎么裁 / 题目说明不读)。"""
    import openpyxl                                   # 只有读 xlsx 要它
    wb = openpyxl.load_workbook(path, data_only=True)
    want = [sheet_name(q["id"]) for q in fb.llm_questions(bank)] + [SPOT_SHEET]
    missing = [s for s in want if s not in wb.sheetnames]
    if missing:
        raise IngestRefused(f"{path.name}: 少了这些页 {missing} (页有 {wb.sheetnames})")
    out = []
    for name in want:
        ws = wb[name]
        header = [str(c.value or "") for c in ws[1]]
        if "cell_id" not in header or PICK_HEADER not in header:
            raise IngestRefused(f"{path.name}/{name}: 表头没有 cell_id 或「{PICK_HEADER}」")
        ci, pi = header.index("cell_id"), header.index(PICK_HEADER)
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[ci] is None:
                continue
            v = row[pi]
            out.append({"sheet": name, "cell_id": str(row[ci]), "value": None if v is None else str(v).strip()})
    return out


def validate_rulings(entries: list[dict], cells: list[dict], bank: dict) -> tuple[dict[str, str], dict, list[str]]:
    """→ ({cell_id: 答案}, 计数, 错误)。错误非空 = 整份不出。"""
    idx = fb.question_index(bank)
    by_id = {c["cell_id"]: c for c in cells}
    label_to_q = {SHORT[q["id"]]: q["id"] for q in fb.llm_questions(bank)}
    errs: list[str] = []
    seen: set[str] = set()
    accepted: dict[str, str] = {}
    counts = {"rows": len(entries), "accepted": 0, "unsure": 0, "blank": 0}
    for e in entries:
        cid = e["cell_id"]
        if cid in seen:
            errs.append(f"{cid}: 出现了两次")
            continue
        seen.add(cid)
        c = by_id.get(cid)
        if c is None:
            errs.append(f"{cid}: 不在格子表里 (表不是这份格子表做的?)")
            continue
        if e["sheet"] == SPOT_SHEET:
            if c["kind"] != "spot":
                errs.append(f"{cid}: 在「{SPOT_SHEET}」页, 格子表里却是 {c['kind']}")
                continue
        elif label_to_q.get(e["sheet"]) != c["question_id"] or c["kind"] != "disagree":
            errs.append(f"{cid}: 在「{e['sheet']}」页, 格子表里是 {c['question_id']} / {c['kind']}")
            continue
        v = e["value"]
        if v in (None, ""):
            counts["blank"] += 1
            continue
        if v == UNSURE:
            counts["unsure"] += 1
            continue
        if v not in fb.closed_set(idx[c["question_id"]]):
            errs.append(f"{cid}: 「{v}」不在 {c['question_id']} 的闭集 {fb.closed_set(idx[c['question_id']])}")
            continue
        accepted[cid] = v
    counts["accepted"] = len(accepted)
    return accepted, counts, errs


def check_owner_run_tag(bank: dict, run_tag: str) -> None:
    err = fb.check_run_tag(bank, run_tag)
    if err:
        raise IngestRefused(err)
    model_tags = {s.split("=", 1)[1].split("@", 1)[0] for s in DEFAULT_RUNS}
    if run_tag in model_tags:
        raise IngestRefused(f"run_tag {run_tag!r} 是模型的一趟; owner 的裁决落 {OWNER_RUN_TAG}")


def build_rows(accepted: dict[str, str], cells: list[dict], bank: dict, run_tag: str) -> list[dict]:
    check_owner_run_tag(bank, run_tag)
    idx = fb.question_index(bank)
    by_id = {c["cell_id"]: c for c in cells}
    now = _iso_now()
    rows = []
    for cid in sorted(accepted):
        c = by_id[cid]
        rows.append({"subject_type": "note", "subject_id": c["note_id"], "question_id": c["question_id"],
                     "question_version": int(idx[c["question_id"]]["version"]),
                     "bank_version": bank["bank_version"], "bank_sha256": bank["_sha256"],
                     "extractor": OWNER_EXTRACTOR, "run_tag": run_tag, "answer": accepted[cid],
                     "evidence": None, "prob": None, "invalid_reason": None, "extracted_at": now})
    return rows


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--file", type=Path, required=True, help="owner 填回的 xlsx")
    ap.add_argument("--cells", type=Path, required=True, help="做表用的那份格子表")
    ap.add_argument("--bank", default="reader")
    ap.add_argument("--run-tag", default=OWNER_RUN_TAG)
    ap.add_argument("--sql-out", type=Path, help="把行写成 upsert SQL (交给有写权限的人执行)")
    a = ap.parse_args(argv)
    try:
        bank = load_secondary(a.bank)
        check_owner_run_tag(bank, a.run_tag)
        cells = read_cells(a.cells)
        accepted, counts, errs = validate_rulings(read_review(a.file, bank), cells, bank)
        if errs:
            for e in errs[:20]:
                print(f"  - {e}", file=sys.stderr)
            raise IngestRefused(f"{len(errs)} 处校验不过, 整份不出")
        rows = build_rows(accepted, cells, bank, a.run_tag)
    except (IngestRefused, Gate1Refused) as exc:
        print(f"拒绝: {exc}", file=sys.stderr)
        return 2
    print(f"{a.file.name}: {counts['rows']} 格 · 收 {counts['accepted']} · 拿不准 {counts['unsure']} · 空着 {counts['blank']} "
          f"→ {len(rows)} 行 ({OWNER_EXTRACTOR} / {a.run_tag} / {bank['bank_version']} {bank['_sha256'][:12]})")
    if a.sql_out:
        sql = rows_to_sql(rows)
        a.sql_out.write_text(sql, encoding="utf-8")
        print(f"SQL 写到 {a.sql_out} ({len(sql.encode('utf-8'))} 字节)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
