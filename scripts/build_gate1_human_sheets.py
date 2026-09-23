#!/usr/bin/env python3
"""闸一人工标注表生成器 (docs/28 §6.1)。

输入: 按固定种子抽好的 100 篇 (5 个项目 × 10 爆 + 10 趴, v_l2_labels, 有正文),
      每个项目一个 JSON: [{note_id, project_id, y, tier, rk, md5, raw_content}, ...]。
      抽样 SQL 见 data-analysis/gate1-human-annotation-plan-2026-09-28.md。
输出: 三份互相看不见的 xlsx (标注员 A / B / C 各一份) + 一份名单 CSV (存档, 回收对账用)。

设计 (owner 2026-09-22:「做互相看不见的表格, 尽量简单一点, 大家都很忙」):
  · 每人 50 篇。100 篇分 4 组各 25 篇 (每组 = 5 个项目各 5 篇):
        A = 组1 + 组2   B = 组2 + 组3   C = 组4 + 组1
    → 组1(A,C)、组2(A,B) 被两人各标一遍 = 50 篇双标 (docs/28 §6.1 的要求),
      组3、组4 各一人标。总工作量 150 篇次, 与原计划「100 篇 + 50 篇再标一遍」相同。
  · 人和模型看同样的东西: 只给代码切好的标题 / 正文 (build_spans, 正文截到 1500 字),
    不给项目名、产品名、爆没爆。模型本来就不问的格子 (没标题 / 正文太短) 预先灰掉写「—」。
  · 每人的顺序用 md5(note_id + 标注员) 单独打乱, 两个人标同一篇时前后文不同。

挡不住什么 (D-051):
  · 挡不住标注员互相商量 —— 表格分开只能让"互相看到"变难, 不能让它不发生。
  · 5 道正例稀少的题 (如 divisive_claim) 在 25 篇重叠里可能全是「否」, 那时 κ 算不出来,
    回收报告要写「不可算」, 不能写成「一致」。
  · 抽样只覆盖闸一的 5 个项目, 不代表全库。

跑法: python build_gate1_human_sheets.py --in-dir <抽样 JSON 目录> --out-dir <输出目录>
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

import feature_bank as fb

ROOT = Path(__file__).resolve().parent.parent
PROJECTS = ("NUC_phase1", "NRT_phase2", "NRT_phase3", "OKMAN_phase1", "SPX_phase1")
PER_PROJECT_Y = 10                         # 每个项目每个 y 抽 10 篇
BLOCK_OF = {"A": (1, 2), "B": (2, 3), "C": (4, 1)}
SEED_TAG = "gate1-human-20260928"          # 与抽样 SQL 里的种子同一个字串

from gate1_labels import SHORT  # 短标签的唯一来源 (零依赖模块; ingest / agreement 也从那儿拿)
SCOPE_ZH = {
    "title": "只看标题",
    "first_sentence": "只看正文第一句（跳过开头的话题标签和表情）",
    "last_para": "只看正文最后一段（不算结尾的 #话题 和 @）",
    "body": "看正文",
    "full": "看标题 + 正文",
}

FONT = "Arial"
F_BODY = Font(name=FONT, size=10)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_GREY = Font(name=FONT, size=10, color="808080")
FILL_HEAD = PatternFill("solid", fgColor="44546A")
FILL_SKIP = PatternFill("solid", fgColor="D9D9D9")
FILL_TODO = PatternFill("solid", fgColor="FFFFFF")
FILL_EX = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
SKIP = "—"


def _h(s: str) -> str:
    return hashlib.md5(s.encode("utf-8")).hexdigest()


def load_sample(in_dir: Path) -> list[dict]:
    rows: list[dict] = []
    for p in PROJECTS:
        d = json.loads((in_dir / f"sample_{p}.json").read_text(encoding="utf-8"))
        bad = [r["note_id"] for r in d if _h(r["raw_content"]) != r["md5"]]
        if bad:
            raise SystemExit(f"{p}: 原文 md5 对不上 {bad} —— 抄错了, 不许生成")
        ys = sorted({r["y"] for r in d})
        cnt = {y: sum(1 for r in d if r["y"] == y) for y in ys}
        if cnt != {0: PER_PROJECT_Y, 1: PER_PROJECT_Y}:
            raise SystemExit(f"{p}: 期望 y=0/1 各 {PER_PROJECT_Y} 篇, 实际 {cnt}")
        rows.extend(d)
    return rows


def assign_blocks(rows: list[dict]) -> None:
    """组号 = (rk-1 + 偏移) % 4 + 1。y=0 偏移 2 格, 于是每组正好每个项目 5 篇。"""
    for r in rows:
        off = 0 if r["y"] == 1 else 2
        r["block"] = (int(r["rk"]) - 1 + off) % 4 + 1
    for p in PROJECTS:
        per = [sum(1 for r in rows if r["project_id"] == p and r["block"] == b) for b in (1, 2, 3, 4)]
        if per != [5, 5, 5, 5]:
            raise SystemExit(f"{p}: 分组不均 {per}")


def add_spans(rows: list[dict], bank: dict) -> None:
    modes = {p: (yaml.safe_load((ROOT / "mappings" / f"{p}.yaml").read_text(encoding="utf-8"))
                 .get("title_extraction", "none")) for p in PROJECTS}
    for r in rows:
        sp = fb.build_spans(r["raw_content"], mode=modes[r["project_id"]])
        r["title"] = sp["title"]
        r["body"] = sp["body"] + ("\n\n……（后面截掉了，模型也只看到这里）" if sp["_truncated"] else "")
        r["title_how"] = sp["_title_how"]
        r["truncated"] = sp["_truncated"]
        skipped = []
        for q in fb.llm_questions(bank):
            sc = q["scope"]
            if sc == "title" and sp["title"] is None:
                skipped.append(q["id"])
            elif fb.visible_len(sp[sc] or "") < (fb.MIN_BODY_CHARS if sc in ("body", "full") else 2):
                skipped.append(q["id"])
        r["skipped"] = skipped


def _as_list(v) -> list[str]:
    if v is None:
        return []
    return [str(x) for x in v] if isinstance(v, list) else [str(v)]


def _header_comment(i: int, q: dict) -> str:
    lines = [f"第 {i} 题：{q['ask']}", f"（{SCOPE_ZH[q['scope']]}）"]
    if q["type"] == "choice":
        for o in q.get("options") or []:
            lines.append(f"· {o['value']}：{o.get('means', '')}")
    else:
        if q.get("yes_if"):
            lines.append(f"算「是」：{q['yes_if']}")
        if q.get("no_if"):
            lines.append(f"算「否」：{q['no_if']}")
    lines.append("更多例子见「题目说明」那一页。")
    return "\n".join(lines)


def _row_height(title: str | None, body: str) -> float:
    # 正文列宽 62 ≈ 每行 30 个汉字; 每行 13.5pt; Excel 行高上限 409。
    lines = sum(max(1, -(-len(seg) // 30)) for seg in body.split("\n"))
    lines = max(lines, -(-len(title or "") // 9) if title else 1)
    return float(min(409, max(60, lines * 13.5 + 8)))


def build_workbook(person: str, rows: list[dict], bank: dict, out: Path) -> None:
    qs = fb.llm_questions(bank)
    mine = [r for r in rows if r["block"] in BLOCK_OF[person]]
    mine.sort(key=lambda r: _h(f"{r['note_id']}:{person}:{SEED_TAG}"))
    n = len(mine)

    wb = Workbook()

    # ── 页 1: 怎么标 ───────────────────────────────────────────────
    ws0 = wb.active
    ws0.title = "怎么标"
    ws0.sheet_view.showGridLines = False
    ws0.column_dimensions["A"].width = 4
    ws0.column_dimensions["B"].width = 100
    ws0["B1"] = f"笔记标注 · 标注员 {person} · 共 {n} 篇"
    ws0["B1"].font = F_TITLE
    steps = [
        "1. 在「标注」那一页，每一行是一篇笔记。先读标题和正文，再往右一题一题选。",
        "2. 点一下白色格子会出现下拉，选「是 / 否」，或者选一个选项。拿不准就选最接近的，不要空着。",
        "3. 灰色写着「—」的格子不用填（那篇没有标题，或者正文太短，这道题本来就不问）。",
        "4. 题目看不懂：鼠标停在表头上会弹出说明；也可以看「题目说明」那一页，每题都有算和不算的例子。",
        "5. 请自己独立标，不要和别人商量，也不要看别人的表。这是整件事唯一的硬要求。",
        f"6. 三天标完就行，每天 {-(-n // 3)} 篇左右、一个小时上下。标完把这个文件直接发回。",
        "7. 最右边「拿不准的地方」可以不填；想说一句就写一句。",
    ]
    for i, s in enumerate(steps, start=3):
        ws0[f"B{i}"] = s
        ws0[f"B{i}"].font = F_BODY
        ws0[f"B{i}"].alignment = WRAP_TOP
    # 不放任何公式: 这个沙箱里 LibreOffice 连三格的表都算不动, 没法重算校验;
    # 而进度条只是锦上添花 —— 白格子就是待办, 一眼能看出还剩多少 (owner: 尽量简单)。
    ws0["B11"] = f"共 {n} 篇 × 20 题。白色格子是要填的，灰色「—」不用填。"
    ws0["B11"].font = F_BOLD

    # 示例行 (只示范格式, 不是任何一篇真实笔记)
    ws0["B15"] = "示例：填好的一行大概长这样（这是编的，不在你要标的笔记里）"
    ws0["B15"].font = F_BOLD
    ex = ws0.cell(row=16, column=2)
    ex.value = ("标题：戒烟第17天，同事都问我怎么忍住的？\n"
                "正文：上周五加班到十点，隔壁工位的老王递过来一根烟，我手都伸出去了又缩回来。"
                "戒烟第17天，最难的不是想抽，是饭后那十分钟不知道手往哪放。你们戒的时候最难熬的是哪一段？\n"
                "→ 标题是问句：是 · 第一句类型：具体事件 · 具体时间：是 · 具体地点或场合：是 · "
                "别人说的原话：否 · 结尾问读者：是 · 亲身经历：是 · 产品角色：未出现 ……")
    ex.font = F_GREY
    ex.fill = FILL_EX
    ex.alignment = WRAP_TOP
    ws0.row_dimensions[16].height = 80

    # ── 页 2: 标注 ─────────────────────────────────────────────────
    ws = wb.create_sheet("标注")
    heads = ["序号", "标题", "正文"] + [f"{i}\n{SHORT[q['id']]}" for i, q in enumerate(qs, 1)] + ["拿不准的地方\n（可不填）", "note_id"]
    for c, h in enumerate(heads, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = F_HEAD
        cell.fill = FILL_HEAD
        cell.alignment = CENTER
        cell.border = BOX
    for i, q in enumerate(qs, 1):
        cm = Comment(_header_comment(i, q), "说明")
        cm.width, cm.height = 380, 220
        ws.cell(row=1, column=3 + i).comment = cm
    ws.row_dimensions[1].height = 48
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 62
    for i in range(len(qs)):
        ws.column_dimensions[_col(4 + i)].width = 9.5
    ws.column_dimensions[_col(4 + len(qs))].width = 24
    ws.column_dimensions[_col(5 + len(qs))].hidden = True
    ws.freeze_panes = "D2"

    dv_bool = DataValidation(type="list", formula1='"是,否"', allow_blank=True,
                             showErrorMessage=True, errorTitle="请从下拉里选", error="只能选「是」或「否」")
    ws.add_data_validation(dv_bool)
    dv_choice = {}
    for q in qs:
        if q["type"] == "choice":
            vals = ",".join(o["value"] for o in q["options"])
            dv = DataValidation(type="list", formula1=f'"{vals}"', allow_blank=True,
                                showErrorMessage=True, errorTitle="请从下拉里选", error="请从下拉选项里选一个")
            ws.add_data_validation(dv)
            dv_choice[q["id"]] = dv

    for ri, r in enumerate(mine, start=2):
        ws.cell(row=ri, column=1, value=ri - 1).alignment = CENTER
        ws.cell(row=ri, column=2, value=r["title"] or "（没有标题）").alignment = WRAP_TOP
        ws.cell(row=ri, column=3, value=r["body"]).alignment = WRAP_TOP
        for c in (1, 2, 3):
            ws.cell(row=ri, column=c).font = F_BODY
            ws.cell(row=ri, column=c).border = BOX
        for qi, q in enumerate(qs):
            cell = ws.cell(row=ri, column=4 + qi)
            cell.border = BOX
            cell.alignment = CENTER
            cell.font = F_BODY
            if q["id"] in r["skipped"]:
                cell.value = SKIP
                cell.fill = FILL_SKIP
                cell.font = F_GREY
            else:
                cell.fill = FILL_TODO
                (dv_choice.get(q["id"]) or dv_bool).add(cell.coordinate)
        note = ws.cell(row=ri, column=4 + len(qs))
        note.alignment = WRAP_TOP
        note.border = BOX
        note.font = F_BODY
        ws.cell(row=ri, column=5 + len(qs), value=r["note_id"])
        ws.row_dimensions[ri].height = _row_height(r["title"], r["body"])

    # ── 页 3: 题目说明 ─────────────────────────────────────────────
    w3 = wb.create_sheet("题目说明")
    h3 = ["题号", "简称", "完整问题", "看哪一段", "算「是」/ 选项说明", "算「否」", "「是」的例子", "「否」的例子"]
    for c, h in enumerate(h3, 1):
        cell = w3.cell(row=1, column=c, value=h)
        cell.font = F_HEAD
        cell.fill = FILL_HEAD
        cell.alignment = CENTER
    for c, wdt in zip("ABCDEFGH", (5, 14, 30, 22, 46, 34, 34, 30)):
        w3.column_dimensions[c].width = wdt
    for i, q in enumerate(qs, 1):
        if q["type"] == "choice":
            yes = "\n".join(f"· {o['value']}：{o.get('means', '')}（例：{o.get('example', '')}）"
                            for o in q.get("options") or [])
            no, yex, nex = "（单选题，选最贴近的一项）", "", ""
        else:
            yes = q.get("yes_if", "")
            no = q.get("no_if", "")
            yex = "\n".join(f"· {x}" for x in _as_list(q.get("yes_examples")))
            nex = "\n".join(f"· {x}" for x in _as_list(q.get("no_examples")))
        vals = [i, SHORT[q["id"]], q["ask"], SCOPE_ZH[q["scope"]], yes, no, yex, nex]
        for c, v in enumerate(vals, 1):
            cell = w3.cell(row=1 + i, column=c, value=v)
            cell.font = F_BODY
            cell.alignment = WRAP_TOP
            cell.border = BOX
    w3.freeze_panes = "C2"

    wb.active = 0
    wb.save(out)


def _col(i: int) -> str:
    s = ""
    while i:
        i, rem = divmod(i - 1, 26)
        s = chr(65 + rem) + s
    return s


def write_manifest(rows: list[dict], path: Path) -> None:
    who = {b: "".join(p for p, bs in BLOCK_OF.items() if b in bs) for b in (1, 2, 3, 4)}
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["note_id", "project_id", "y", "tier", "rk", "block", "annotators",
                    "title_how", "body_truncated", "skipped_questions", "raw_md5"])
        for r in sorted(rows, key=lambda r: (r["block"], r["project_id"], -r["y"], r["rk"])):
            w.writerow([r["note_id"], r["project_id"], r["y"], r["tier"], r["rk"], r["block"],
                        who[r["block"]], r["title_how"], int(r["truncated"]),
                        "|".join(r["skipped"]), r["md5"]])


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--in-dir", required=True, type=Path)
    ap.add_argument("--out-dir", required=True, type=Path)
    ap.add_argument("--manifest", type=Path,
                    default=ROOT / "data-analysis" / "gate1-human-sample-2026-09-28.csv")
    a = ap.parse_args()
    bank = fb.load_bank()
    if len(fb.llm_questions(bank)) != 20:
        raise SystemExit(f"题库题数变了: {len(fb.llm_questions(bank))} (本表按 20 题设计)")
    rows = load_sample(a.in_dir)
    assign_blocks(rows)
    add_spans(rows, bank)
    a.out_dir.mkdir(parents=True, exist_ok=True)
    for person in BLOCK_OF:
        build_workbook(person, rows, bank, a.out_dir / f"笔记标注_标注员{person}.xlsx")
    write_manifest(rows, a.manifest)
    print(f"ok: {len(rows)} 篇 → 3 份表 (各 50 篇) + 名单 {a.manifest}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
