#!/usr/bin/env python3
"""check_gate1_review.py —— 副题库闸一的样本 / 裁决表 / 收表 (D-128) 自检。只用合成数据 (加一道对进了仓库的名单的绊线)。

    cd scripts && python check_gate1_review.py           # CI: 纯函数 (不要 openpyxl)
    python check_gate1_review.py --xlsx                  # 本地: 真做一份 xlsx → 填 → 收回 (要 openpyxl)

反证 (改坏之后必须变红):
  ① 抽样: 留出半的篇混进发现半 / note_id 重复 / 分层题多行 / note_id 不像 <项目>_rec / 项目没 mapping / 层不够数 /
     y 不是 0/1 / md5 坏了 → 拒; 层内不按 sha256(种子:note_id) 排、D-128 点名的 8 篇没剔、闸二口径没加    → §1 红
  ② 进了仓库的名单: 不是 200 篇 / 不是 150 + 50 / 有留出半的篇 / 有点名剔掉的篇 / 层内顺序不对 / 层和三道答案对不上 → §2 红
  ③ 裁决表: 正文 md5 对不上还做表 / 带正文的文件写进仓库里没被 ignore 的路径 / 甲乙泄露是哪一趟 / 表里出现模型名  → §3 红
  ④ 收表: cell_id 不在格子表 / 题页上放了别的题或抽查格 / 抽查页上放了不一致格 / 同一格两次 / 答案不在闭集 → 拒;
     拿不准和空着的写了行; run_tag 用 primary / gate1-* / 模型的那一趟; 行不是 human:owner + 副题库版本和 sha;
     owner 的理由或正文进了 SQL                                                                              → §4 红
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import pathlib
import sys
import tempfile

import feature_bank as fb
import gate1_models as gm
import gate1_sample as gs
from discover_feature_candidates import in_discovery_half
from gate1_labels import SHORT

ROOT = pathlib.Path(__file__).resolve().parent.parent
SENTINEL = "哨兵正文SENTINEL-c41e"
BANK = gm.load_secondary("reader")
MANIFEST = ROOT / "data-analysis" / "gate1-reader-sample-2026-10-26.csv"


def _ids(n: int, half: str, prefix: str = "NUC_phase1") -> list[str]:
    out, i = [], 0
    while len(out) < n:
        nid = f"{prefix}_recSyn{i:05d}"
        if in_discovery_half(nid) == (half == "discovery"):
            out.append(nid)
        i += 1
    return out


def _pool(n_a: int, n_b: int, half: str = "discovery") -> list[dict]:
    rows = []
    for k, nid in enumerate(_ids(n_a + n_b + 3, half)):
        a = k < n_a
        rows.append({"note_id": nid, "project_id": "NUC_phase1", "y": k % 2, "ear": "是" if a and k % 3 == 0 else "否",
                     "afh": "是" if a and k % 3 == 1 else "否", "tiq": ("是" if a and k % 3 == 2 else None),
                     "n_ear": 1, "n_afh": 1, "n_tiq": 1, "raw_md5": hashlib.md5(nid.encode()).hexdigest(),
                     "has_essence": True, "body_len": 120, "vanished": False})
    rows[-1].update(afh=None)                      # 两题没答全: 不进池子
    rows[-2].update(ear=None, n_ear=0)
    rows[-3].update(ear="否", afh="否", tiq=None)  # tiq NULL 按「不是」→ B 层
    return rows


def _refused(fn, needle: str, what: str, exc_type=gs.SampleRefused) -> None:
    try:
        fn()
    except exc_type as exc:
        assert needle in str(exc), f"{what}: 拒了但原因不对 (要含 {needle!r}): {exc}"
        return
    raise AssertionError(f"{what}: 该拒没拒")


# §1 抽样 ──────────────────────────────────────────────────────────────────────────
def check_sampler(td: pathlib.Path) -> None:
    pool = _pool(30, 12)
    rows, st = gs.build_sample(pool, take={"A": 10, "B": 5})
    assert st["rows_with_any_answer"] == 45 and st["pool_both_answered"] == 43 and st["pool"] == 43, st
    assert st["pool_by_stratum"] == {"A": 30, "B": 13} and st["taken"] == {"A": 10, "B": 5}, st
    for s in ("A", "B"):
        got = [r["note_id"] for r in rows if r["stratum"] == s]
        want = sorted((r["note_id"] for r in pool if r["ear"] is not None and r["afh"] is not None and gs.stratum(r) == s),
                      key=lambda n: hashlib.sha256(f"{gs.SEED}:{n}".encode()).hexdigest())
        assert got == want[:len(got)], s
    assert [r["stratum"] for r in rows] == ["A"] * 10 + ["B"] * 5
    rows2, _ = gs.build_sample(pool, take={"A": 10, "B": 5}, seed="other-seed")
    assert [r["note_id"] for r in rows2] != [r["note_id"] for r in rows], "换种子必须换样本 (种子真的进了排序)"
    # D-128 点名剔掉的篇
    ex = next(n for n in gs.EXCLUDE if in_discovery_half(n))
    p2 = [dict(r) for r in pool]
    p2[0].update(note_id=ex, project_id=ex.split("_rec")[0])
    rows3, st2 = gs.build_sample(p2, take=None)
    assert st2["excluded_present"] == [ex] and st2["pool"] == 42 and ex not in {r["note_id"] for r in rows3}, st2
    # 拒
    bad = [dict(r) for r in pool] + [dict(pool[0], note_id=_ids(1, "holdout")[0])]
    _refused(lambda: gs.build_sample(bad), "不在发现半", "留出半的篇混进来")
    _refused(lambda: gs.build_sample(pool + [dict(pool[0])]), "重复", "note_id 重复")
    _refused(lambda: gs.build_sample([dict(pool[0], n_afh=2)] + pool[1:]), "多行", "分层题多行")
    _refused(lambda: gs.build_sample([dict(pool[0], project_id="NRT_phase2")] + pool[1:]), "<project_id>_rec", "项目对不上")
    _refused(lambda: gs.build_sample(pool, mappings_dir=td), "没有 mapping", "项目没 mapping")
    _refused(lambda: gs.build_sample(pool, take={"A": 31, "B": 5}), "不够数", "层不够数")
    assert gs.parse_take("", "discovery") == {"A": 150, "B": 50} and gs.parse_take("", "holdout") is None
    _refused(lambda: gs.build_sample(pool, take=gs.parse_take("", "discovery")), "不够数", "默认 150 / 50 不够数")
    _refused(lambda: gs.build_sample([dict(pool[0], y=2)] + pool[1:]), "不是 0/1", "y 坏了")
    _refused(lambda: gs.build_sample([dict(pool[0], raw_md5="x")] + pool[1:]), "raw_md5", "md5 坏了")
    # 留出半: 闸二口径 (essence / 正文 ≥ 50 / 没删); body_len 未知的不剔
    hp = _pool(8, 4, "holdout")
    hp[0]["vanished"] = True
    hp[1]["has_essence"] = False
    hp[2]["body_len"] = 49
    hp[3]["body_len"] = None
    hrows, hst = gs.build_sample(hp, half="holdout", take=None)
    assert hst["gate2_filter_drops"] == {"vanished": 1, "no_essence": 1, "body_lt_50": 1} and len(hrows) == 12 - 3 + 1, hst
    _refused(lambda: gs.build_sample(hp, half="discovery"), "不在发现半", "留出半当发现半抽")
    # SQL: 快照 sha / Sonnet / 版本 / 半边
    sql = gs.pool_sql("discovery")
    snap = BANK["gate2"]["condition_snapshot"]
    assert f"bank_sha256 = '{snap['bank_sha256']}'" in sql and "extractor = 'llm:claude-sonnet-5-5'" in sql
    assert "run_tag = 'primary'" in sql and "question_version = 1" in sql and "'0', '2'" in sql and "'1'" not in sql
    assert "'1', '3'" in gs.pool_sql("holdout") and "evidence" not in sql and "raw_content, ''))" in sql
    # CSV 往返 + --verify
    p = td / "m.csv"
    p.write_text(gs.to_csv(rows), encoding="utf-8")
    back = gs.read_manifest(p)
    assert [r["note_id"] for r in back] == [r["note_id"] for r in rows] and back[0]["y"] == rows[0]["y"]
    assert any(r["tiq"] is None for r in gs.read_manifest(p)) or all(r["tiq"] for r in rows)
    j = td / "pool.json"
    j.write_text(json.dumps([[r[c] for c in gs.POOL_COLUMNS] for r in pool], ensure_ascii=False), encoding="utf-8")
    import contextlib
    import io
    with contextlib.redirect_stdout(io.StringIO()):
        assert gs.main(["--from-json", str(j), "--take", "10,5", "--verify", str(p)]) == 0
        p.write_text(gs.to_csv(rows[:-1]), encoding="utf-8")
        assert gs.main(["--from-json", str(j), "--take", "10,5", "--verify", str(p)]) == 1
    with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
        assert gs.main(["--from-json", str(j)]) == 2          # 默认 150/50, 合成池子不够 → 拒, 退出 2
    print("  §1 抽样: 层内按 sha256(种子:note_id)、点名的篇剔掉、两题没答全不进池、tiq NULL 算「不是」; 留出半 / 重复 / 多行 / "
          "项目 / mapping / 不够数 / y / md5 → 拒; 留出半加闸二口径; SQL 钉快照; --verify ✓")


# §2 进了仓库的名单 (绊线: 内部自洽; 池子完整性要连库才查得了, 那是 gate1_sample --verify 的事) ─────────────────
def check_committed_manifest() -> None:
    rows = gs.read_manifest(MANIFEST)
    assert len(rows) == 200 and [r["stratum"] for r in rows] == ["A"] * 150 + ["B"] * 50, "D-128: A 150 + B 50"
    assert all(in_discovery_half(r["note_id"]) for r in rows), "名单里有留出半的篇 —— 闸二会被污染"
    assert not {r["note_id"] for r in rows} & set(gs.EXCLUDE), "D-128 点名剔掉的篇进了名单"
    for s in ("A", "B"):
        ks = [gs.rank_key(r["note_id"]) for r in rows if r["stratum"] == s]
        assert ks == sorted(ks), f"{s} 层不是按 sha256({gs.SEED}:note_id) 升序取的"
    for r in rows:
        assert gs.stratum(r) == r["stratum"], r["note_id"]
        assert r["ear"] is not None and r["afh"] is not None, r["note_id"]
        assert r["note_id"].startswith(r["project_id"] + "_rec") and (ROOT / "mappings" / f"{r['project_id']}.yaml").is_file()
    text = MANIFEST.read_text(encoding="utf-8")
    assert text.splitlines()[0] == ",".join(gs.OUT_COLUMNS) and "raw_content" not in text
    sha = hashlib.sha256(text.encode("utf-8")).hexdigest()
    print(f"  §2 名单 {MANIFEST.name}: 200 篇 = A 150 + B 50, 都在发现半, 没有点名剔掉的篇, 层内按种子排序, 层与答案自洽 "
          f"(sha256 {sha[:12]}) ✓")


# §3 裁决表 (纯函数) ────────────────────────────────────────────────────────────────
def _synthetic(td: pathlib.Path):
    import build_gate1_review_sheet as br
    notes = [f"NUC_phase1_recRev{i:03d}" for i in range(30)]
    texts = []
    for i, n in enumerate(notes):
        raw = f"【标题】早饭吃什么第{i}天\n【正文】这周每天七点起来做早饭，燕麦配鸡蛋。{SENTINEL}\n\n你们早饭一般吃什么？"
        texts.append({"note_id": n, "project_id": "NUC_phase1", "title": None, "raw_content": raw})
    man = td / "rev_manifest.csv"
    with open(man, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(gs.OUT_COLUMNS)
        for t in texts:
            w.writerow([t["note_id"], "NUC_phase1", 0, "是", "否", "否", "A", hashlib.md5(t["raw_content"].encode()).hexdigest()])
    cells = []
    for i, n in enumerate(notes):
        for q in ("reader_ask_form", "invites_sharing"):
            if i < 12:
                kind = "disagree"
                a, b, h = ("挑一个", "说自己的", "求解答") if q == "reader_ask_form" else ("是", "否", "是")
            elif i < 20:
                kind = "spot"
                a, b, h = ("说自己的",) * 3 if q == "reader_ask_form" else ("否",) * 3
            else:
                kind = "agree"
                a, b, h = ("没问读者",) * 3 if q == "reader_ask_form" else ("否",) * 3
            if kind == "disagree" and i % 2:
                a, b = b, a                                       # 甲乙不该跟着哪一趟走
            cells.append({"cell_id": f"{n}:{q}", "note_id": n, "question_id": q, "stratum": "A", "kind": kind,
                          "A": a, "B": b, "H": h, "A_reason": None, "B_reason": None, "H_reason": None,
                          "owner": None, "reference": None})
    cp = td / "rev_cells.csv"
    cp.write_text(gm.cells_csv(cells), encoding="utf-8")
    tp = td / "texts.json"
    tp.write_text(json.dumps([{"pivot": json.dumps(texts, ensure_ascii=False)}], ensure_ascii=False), encoding="utf-8")
    return br, gm.read_cells(cp), gs.read_manifest(man), tp, texts, cp, man


def check_review_plan(td: pathlib.Path) -> None:
    br, cells, manifest, tp, texts, _, _ = _synthetic(td)
    needed = {c["note_id"] for c in cells if c["kind"] in ("disagree", "spot")}
    got = br.load_texts(tp, needed, manifest)
    plan = br.plan_review(cells, BANK, got, manifest)
    raf = plan["pages"]["reader_ask_form"]
    assert len(raf) == 12 and len(plan["pages"]["invites_sharing"]) == 12 and len(plan["spot"]) == 16
    assert all(r["options"][-1] == br.UNSURE and r["options"][:-1] == list(fb.closed_set(fb.question_index(BANK)[r["question_id"]]))
               for r in raf + plan["spot"])
    # 甲 / 乙 / 丙 = 不同的答案, 顺序和 A/B/H 无关: 「挑一个」(有的格是 A 的, 有的是 B 的) 既当过甲也当过别的
    marks_of = {m for r in raf for m, v in r["candidates"] if v == "挑一个"}
    assert {tuple(v for _, v in r["candidates"]) for r in raf} and len(marks_of) >= 2, marks_of
    assert all(sorted(v for _, v in r["candidates"]) == sorted(["挑一个", "说自己的", "求解答"]) for r in raf)
    flat = json.dumps(plan, ensure_ascii=False)
    for leak in ("sonnet", "haiku", "claude", "llm:", "fqr-g1", '"A"', '"B"', '"H"', "A_reason"):
        assert leak not in flat.lower() if leak.islower() else leak not in flat, f"表里漏了 {leak!r}"
    assert all(r["model_answer"] in ("说自己的", "否") for r in plan["spot"])
    # 正文 md5 对不上 → 拒
    t2 = [dict(t) for t in texts]
    t2[0]["raw_content"] += "改了一个字"
    bad = td / "texts_bad.json"
    bad.write_text(json.dumps(t2, ensure_ascii=False), encoding="utf-8")
    _refused(lambda: br.load_texts(bad, needed, manifest), "md5", "正文改过还做表", br.ReviewRefused)
    _refused(lambda: br.load_texts(tp, needed | {"NUC_phase1_recNope"}, manifest), "缺", "正文缺篇", br.ReviewRefused)
    # 路径: 仓库里没被 ignore → 拒; 仓库里被 ignore (*.xlsx) / 仓库外 → 放行
    _refused(lambda: br.check_private_path(ROOT / "scripts" / "owner_texts.json", "正文 JSON"), ".gitignore", "正文写进仓库",
             br.ReviewRefused)
    br.check_private_path(ROOT / "闸一裁决.xlsx", "裁决表")
    br.check_private_path(td / "x.xlsx", "裁决表")
    _refused(lambda: br.check_private_path(ROOT / "data-analysis" / "gate1-reader-cells-x.csv", "带正文的表"), ".gitignore",
             "带正文的文件写成名单例外的路径", br.ReviewRefused)
    sql = br.texts_sql(sorted(needed))
    assert "raw_content" in sql and sql.count("'NUC_phase1_recRev") == len(needed)
    print("  §3 裁决表: 只放不一致 / 抽查格; 甲乙丙 = 不同答案、顺序与哪一趟无关、没有模型名; 下拉 = 全部选项 + 拿不准; "
          "正文 md5 不对 / 缺篇 → 拒; 带正文的文件进仓库 (没被 ignore) → 拒 ✓")


# §4 收表 (纯函数) ──────────────────────────────────────────────────────────────────
def check_ingest(td: pathlib.Path) -> None:
    import ingest_gate1_review as ig
    from build_gate1_review_sheet import SPOT_SHEET, UNSURE
    from ingest_gate1_answers import rows_to_sql
    _, cells, _, _, _, _, _ = _synthetic(td)
    by = {c["cell_id"]: c for c in cells}
    dis = [c for c in cells if c["kind"] == "disagree"]
    spot = [c for c in cells if c["kind"] == "spot"]
    entries = [{"sheet": SHORT[c["question_id"]], "cell_id": c["cell_id"], "value": "求推荐" if c["question_id"] == "reader_ask_form" else "否",
                "reason": SENTINEL} for c in dis]
    entries += [{"sheet": SPOT_SHEET, "cell_id": c["cell_id"], "value": c["A"], "reason": SENTINEL} for c in spot]
    entries[0]["value"] = UNSURE
    entries[1]["value"] = None
    acc, counts, errs = ig.validate_rulings(entries, cells, BANK)
    assert not errs and counts == {"rows": len(entries), "accepted": len(entries) - 2, "unsure": 1, "blank": 1}, (errs, counts)
    rows = ig.build_rows(acc, cells, BANK, gm.OWNER_RUN_TAG)
    assert len(rows) == len(entries) - 2
    assert {(r["extractor"], r["run_tag"], r["bank_version"], r["bank_sha256"]) for r in rows} == {
        ("human:owner", "fqr-g1-owner", "fq-v0.2-cand-reader", BANK["_sha256"])}
    assert {(r["question_id"], r["question_version"]) for r in rows} == {("reader_ask_form", 1), ("invites_sharing", 2)}
    assert all(r["evidence"] is None and r["prob"] is None for r in rows)
    assert all(by[f"{r['subject_id']}:{r['question_id']}"]["kind"] in ("disagree", "spot") for r in rows)
    sql = rows_to_sql(rows)
    assert SENTINEL not in sql and "human:owner" in sql and "fqr-g1-owner" in sql and BANK["_sha256"] in sql
    # 拒
    for bad, needle in (([dict(entries[2], cell_id="NUC_phase1_recNope:reader_ask_form")], "不在格子表"),
                        ([dict(entries[2], sheet=SHORT["invites_sharing"])], "页, 格子表里是"),
                        ([dict(entries[-1], sheet=SHORT[spot[-1]["question_id"]])], "页, 格子表里是"),
                        ([dict(entries[2], sheet=SPOT_SHEET)], "格子表里却是 disagree"),
                        ([entries[2], dict(entries[2])], "出现了两次"),
                        ([dict(entries[2], value="外星选项")], "不在 reader_ask_form 的闭集"),
                        ([dict(entries[2], value="是")], "不在 reader_ask_form 的闭集")):
        _, _, errs = ig.validate_rulings(bad, cells, BANK)
        assert any(needle in e for e in errs), (needle, errs)
    for tag in ("primary", "gate1-20260928", "fqr-g1-sonnet55-a", "fqr-g1-haiku55", "owner", "fqr-"):
        _refused(lambda: ig.build_rows(acc, cells, BANK, tag), "", f"owner run_tag={tag}", ig.IngestRefused)
    print("  §4 收表: 拿不准 / 空着不写行; 行 = human:owner + fqr-g1-owner + 副题库版本与 sha, 没有 evidence; 理由不进 SQL; "
          "cell_id / 页 / 重复 / 闭集 → 拒; primary / gate1-* / 模型那几趟 / 光前缀 → 拒 ✓")


# §5 xlsx 往返 (--xlsx, 要 openpyxl) ──────────────────────────────────────────────────
def check_xlsx_roundtrip(td: pathlib.Path) -> None:
    import openpyxl
    import ingest_gate1_review as ig
    from build_gate1_review_sheet import SPOT_SHEET
    br, cells, manifest, tp, texts, cp, man = _synthetic(td)
    out = td / "闸一裁决.xlsx"
    assert br.main(["--cells", str(cp), "--manifest", str(man), "--texts", str(tp), "--out", str(out)]) == 0
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["怎么裁", SHORT["reader_ask_form"], SHORT["invites_sharing"], SPOT_SHEET, "题目说明"]
    blob = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None)
    for leak in ("sonnet", "haiku", "claude", "llm:", "fqr-g1", "human:"):
        assert leak not in blob.lower(), f"表里有 {leak!r}"
    assert SENTINEL in blob, "正文本来就该在表里 (表不进仓库)"
    ws = wb[SHORT["reader_ask_form"]]
    dv = ws.data_validations.dataValidation
    assert dv and "拿不准" in dv[0].formula1 and "挑一个" in dv[0].formula1
    assert ws.column_dimensions["J"].hidden
    # owner 填: 题页选「求推荐」/「否」, 第一格拿不准; 抽查页照判读的答案
    for name, val in ((SHORT["reader_ask_form"], "求推荐"), (SHORT["invites_sharing"], "否")):
        w = wb[name]
        for r in range(2, w.max_row + 1):
            w.cell(row=r, column=8, value=("拿不准" if r == 2 else val))
            w.cell(row=r, column=9, value=SENTINEL)
    w = wb[SPOT_SHEET]
    for r in range(2, w.max_row + 1):
        w.cell(row=r, column=7, value=w.cell(row=r, column=6).value)
    filled = td / "闸一裁决_填好.xlsx"
    wb.save(filled)
    sql_out = td / "owner.sql"
    assert ig.main(["--file", str(filled), "--cells", str(cp), "--sql-out", str(sql_out)]) == 0
    sql = sql_out.read_text(encoding="utf-8")
    assert SENTINEL not in sql and sql.count("NUC_phase1_recRev") == 12 + 8 and "fqr-g1-owner" in sql
    assert br.main(["--cells", str(cp), "--manifest", str(man), "--texts", str(tp),
                    "--out", str(ROOT / "scripts" / "x.xlsx.json")]) == 2
    print("  §5 xlsx: 页 / 下拉 / 隐藏 cell_id; 表里没有模型名; 填好收回 → SQL 里没有正文和理由; 输出到仓库里没被 ignore 的路径 → 退出 2 ✓")


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", action="store_true", help="再跑 xlsx 往返 (要 openpyxl)")
    a = ap.parse_args(argv)
    assert SHORT["reader_ask_form"] and len(set(SHORT.values())) == len(SHORT), "短标签要唯一 (收表按标签反查题号)"
    with tempfile.TemporaryDirectory() as td:
        t = pathlib.Path(td)
        check_sampler(t)
        check_committed_manifest()
        check_review_plan(t)
        check_ingest(t)
        if a.xlsx:
            check_xlsx_roundtrip(t)
    print(f"\ncheck_gate1_review: {5 if a.xlsx else 4} 节全过" + ("" if a.xlsx else " (xlsx 往返用 --xlsx, 要 openpyxl)"))
    return 0


if __name__ == "__main__":
    sys.exit(main())
