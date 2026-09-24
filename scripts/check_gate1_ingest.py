"""
check_gate1_ingest.py —— 闸一名单的灰格往返 + prob 口径 (D-085)。

    cd scripts && python check_gate1_ingest.py           # 名单级往返 (CI 跑这个; 不需要 openpyxl)
    cd scripts && python check_gate1_ingest.py --xlsx    # 再加一遍真 xlsx 往返 (要 openpyxl; 本地收表的机器上跑)

起因: build_gate1_human_sheets.write_manifest 用 "|" 拼 skipped_questions, ingest_gate1_answers.validate
用 "," 拆。只有一个灰格的篇碰巧没事; NUC_phase1_recv46LaDAdFFc 有 11 个灰格, 拆出来是一个认不出的长串,
灰格校验在两个方向上都失效 —— Jev 在 C 表这 11 格里填的答案原样入了库 (C 表 1000 行, 应为 989)。
CI 以前对这两个脚本只做 py_compile。

守得住什么:
  · join_skipped / split_skipped 往返 (多个灰格、旧的 "," 写法、空格、重复); 生成器的 write_manifest
    必须走 join_skipped (真调它写一份名单, 再用收表脚本的 load_manifest 读回来)。
  · 仓里那份真名单 (data-analysis/gate1-human-sample-2026-09-28.csv) 上: C 表 50 篇灰格照填「—」→ 过,
    989 行; 在 NUC 那 11 个灰格里填答案 (当时发生的事) → 11 条错、整份不写; 「—」出现在非灰格 → 错。
  · 名单拆出题库里没有的题号 → 直接红 (反证: 换回旧的按 "," 拆, 这一条接住它)。
  · prob = 所选答案的概率: Jev 写「判「是」的概率 0.39」、表里答「否」→ 存 0.61; 选择题「第一名只有」照存。
  · --xlsx: build_workbook 出一张带 11 个灰格的表 → 填白格 → read_sheet / validate / build_rows 走一遍。
挡不住什么:
  · 标注员 / Jev 在白格里答错; 名单本身的灰格判据与 feature_bank.render_call 漂开 (那是生成器的事)。
"""

from __future__ import annotations

import argparse
import ast
import csv
import sys
import tempfile
import types
from pathlib import Path

import feature_bank as fb
import gate1_labels as gl
import ingest_gate1_answers as ing

ROOT = Path(__file__).resolve().parent.parent
NUC = "NUC_phase1_recv46LaDAdFFc"


def _import_builder():
    """build_gate1_human_sheets 顶层 import openpyxl。CI 不装它; write_manifest 只用 csv, 模块顶层的
    openpyxl 名字只被当构造器调 (Font(...) 之类) —— 没装时给一副哑壳, 让【真的】write_manifest 能跑。"""
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        class _Any:
            def __init__(self, *a, **k):
                pass
        mods = {n: types.ModuleType(n) for n in ("openpyxl", "openpyxl.comments", "openpyxl.styles",
                                                 "openpyxl.worksheet", "openpyxl.worksheet.datavalidation")}
        mods["openpyxl"].Workbook = _Any
        mods["openpyxl.comments"].Comment = _Any
        for n in ("Alignment", "Border", "Font", "PatternFill", "Side"):
            setattr(mods["openpyxl.styles"], n, _Any)
        mods["openpyxl.worksheet.datavalidation"].DataValidation = _Any
        sys.modules.update(mods)
    import build_gate1_human_sheets as B
    return B


def _answers_for(bank: dict, skipped: set[str], grey=ing.SKIP_MARK) -> dict[str, str]:
    """一篇的 20 格: 灰格填 grey, 白格填闭集里第一个值。"""
    return {q["id"]: (grey if q["id"] in skipped else fb.closed_set(q)[0]) for q in fb.llm_questions(bank)}


def check_split_join() -> None:
    ids = ["opening_type", "has_specific_time", "has_specific_place"]
    cell = gl.join_skipped(ids)
    assert cell == "opening_type|has_specific_time|has_specific_place", cell
    assert gl.split_skipped(cell) == ids
    assert gl.split_skipped("opening_type,has_specific_time") == ids[:2], "旧的逗号写法也要认"
    assert gl.split_skipped(" opening_type | has_specific_time ,opening_type ") == ids[:2], "空格 / 重复"
    assert gl.split_skipped("") == [] and gl.split_skipped(None) == [] and gl.join_skipped([]) == ""
    for bad in (["a|b"], ["a,b"], [""]):
        try:
            gl.join_skipped(bad)
            raise AssertionError(f"join_skipped({bad!r}) 该拒")
        except ValueError:
            pass
    # 生成器必须走 join_skipped —— 源码里 write_manifest 不许再自己拼
    src = (Path(__file__).parent / "build_gate1_human_sheets.py").read_text(encoding="utf-8")
    fn = next(n for n in ast.walk(ast.parse(src)) if isinstance(n, ast.FunctionDef) and n.name == "write_manifest")
    body = ast.get_source_segment(src, fn)
    assert "join_skipped(" in body and '"|".join' not in body and '",".join' not in body, body
    print("  ✓ join/split: 多灰格往返、旧逗号写法、空格重复都对; write_manifest 走 join_skipped")


def check_manifest_roundtrip(bank: dict) -> None:
    B = _import_builder()
    skipped = [q["id"] for q in fb.llm_questions(bank)][:11]
    rows = [
        {"note_id": "X_1", "project_id": "X", "y": 1, "tier": "爆", "rk": 1, "block": 4, "title_how": "colon",
         "truncated": False, "skipped": skipped, "md5": "m1"},
        {"note_id": "X_2", "project_id": "X", "y": 0, "tier": "趴", "rk": 2, "block": 1, "title_how": "colon",
         "truncated": False, "skipped": ["title_is_question"], "md5": "m2"},
        {"note_id": "X_3", "project_id": "X", "y": 0, "tier": "趴", "rk": 3, "block": 4, "title_how": "colon",
         "truncated": False, "skipped": [], "md5": "m3"},
    ]
    with tempfile.TemporaryDirectory() as d:
        path = Path(d) / "m.csv"
        B.write_manifest(rows, path)                   # 生成器的真代码
        man = ing.load_manifest(path)                  # 收表脚本的真代码
    assert gl.split_skipped(man["X_1"]["skipped_questions"]) == skipped
    assert gl.split_skipped(man["X_2"]["skipped_questions"]) == ["title_is_question"]
    assert man["X_3"]["skipped_questions"] == ""
    notes = [{"note_id": n, "answers": _answers_for(bank, set(gl.split_skipped(man[n]["skipped_questions"]))), "probs": {}}
             for n in ("X_1", "X_2", "X_3")]
    errs = ing.validate(Path("C.xlsx"), "C", notes, man, bank)       # C = 组 4 + 组 1 → 三篇都是它的
    assert not errs, errs
    got, blanks = ing.build_rows(notes, bank, "jev:t-C", "gate1-t")
    assert len(got) == 3 * 20 - 11 - 1 and not blanks, len(got)
    print(f"  ✓ write_manifest → load_manifest → validate: 11 格 + 1 格灰格的名单往返后校验通过, {len(got)} 行 (灰格没有行)")


def check_real_manifest(bank: dict) -> None:
    man = ing.load_manifest(ing.DEFAULT_MANIFEST)
    assert NUC in man and len(gl.split_skipped(man[NUC]["skipped_questions"])) == 11, man.get(NUC)
    c_notes = [nid for nid, r in man.items() if "C" in (r["annotators"] or "")]
    assert len(c_notes) == 50

    def sheet(grey_answer: str | None = None, extra_dash: str | None = None) -> list[dict]:
        out = []
        for nid in c_notes:
            sk = set(gl.split_skipped(man[nid]["skipped_questions"]))
            a = _answers_for(bank, sk)
            if grey_answer is not None:
                a.update({q: grey_answer for q in sk})
            if extra_dash and nid == extra_dash:
                white = next(q for q in a if q not in sk)
                a[white] = ing.SKIP_MARK
            out.append({"note_id": nid, "answers": a, "probs": {}})
        return out

    # 灰格照填「—」→ 过; C 表应是 50 × 20 − 11 = 989 行 (当时入库的是 1000)
    ok = sheet()
    assert not ing.validate(Path("C.xlsx"), "C", ok, man, bank)
    rows, _ = ing.build_rows(ok, bank, "jev:1.13.0-C", "gate1-20260928")
    assert len(rows) == 989 and not any(r["subject_id"] == NUC and r["question_id"] in gl.split_skipped(man[NUC]["skipped_questions"]) for r in rows), len(rows)
    # 当时发生的事: Jev 在 11 个灰格里都填了答案 → 现在 11 条错, 整份不写
    errs = ing.validate(Path("C.xlsx"), "C", sheet(grey_answer="否"), man, bank)
    assert len(errs) == 11 and all(NUC in e and "灰格" in e for e in errs), errs
    # 「—」填在非灰格 → 错
    errs = ing.validate(Path("C.xlsx"), "C", sheet(extra_dash=NUC), man, bank)
    assert len(errs) == 1 and "名单没说这题不问" in errs[0], errs
    # 反证: 换回旧的按 "," 拆 → 拆出一个认不出的长串; 新加的「题库里没有的题号」那一条接住它 (整份红)
    old = lambda cell: list(filter(None, (cell or "").split(",")))          # noqa: E731
    assert old(man[NUC]["skipped_questions"])[0].startswith("opening_type|has_specific_time|"), "反证前提不成立"
    saved = ing.split_skipped
    ing.split_skipped = old
    try:
        errs = ing.validate(Path("C.xlsx"), "C", sheet(grey_answer="否"), man, bank)
    finally:
        ing.split_skipped = saved
    assert len(errs) == 1 and "题库没有的题号" in errs[0], errs
    # A 表那篇单灰格 (OKMAN) 照旧: 「—」过, 填答案红
    a_notes = [nid for nid, r in man.items() if "A" in (r["annotators"] or "")]
    a_sheet = [{"note_id": n, "answers": _answers_for(bank, set(gl.split_skipped(man[n]["skipped_questions"]))), "probs": {}} for n in a_notes]
    assert not ing.validate(Path("A.xlsx"), "A", a_sheet, man, bank)
    assert len(ing.build_rows(a_sheet, bank, "jev:1.13.0-A", "gate1-20260928")[0]) == 999
    print("  ✓ 真名单: C 表灰格照填 → 989 行; 11 个灰格填了答案 → 11 条错整份不写; 「—」填错位置 → 错;"
          " 旧的逗号拆法被「题号认不出」接住; A 表单灰格照旧 999 行")


def check_prob(bank: dict) -> None:
    num_to_qid = {4: "has_specific_place", 12: "product_role", 13: "efficacy_promise"}
    text = ("4 具体地点或场合：判「是」的概率 0.39，落在 0.35–0.65\n"
            "12 产品角色：第一名只有 0.45；「主角」0.45 与「只暗示」0.33 太接近\n"
            "13 效果承诺：判「是」的概率 1.7，写错了\n"
            "99 不存在的题：判「是」的概率 0.5")
    probs = ing.parse_probs(text, num_to_qid)
    assert probs == {"has_specific_place": (ing.PROB_YES, 0.39), "product_role": (ing.PROB_TOP, 0.45)}, probs
    assert ing.prob_of_answer("否", probs["has_specific_place"]) == 0.61
    assert ing.prob_of_answer("是", probs["has_specific_place"]) == 0.39
    assert ing.prob_of_answer("主角", probs["product_role"]) == 0.45
    assert ing.prob_of_answer("主角", (ing.PROB_YES, 0.7)) is None, "P(是) 配选择题答案 → 不写"
    assert ing.prob_of_answer("否", None) is None
    notes = [{"note_id": "X_1", "answers": {"has_specific_place": "否", "product_role": "主角", "own_experience": "是"},
              "probs": probs}]
    rows, _ = ing.build_rows(notes, bank, "jev:t-A", "gate1-t")
    by = {r["question_id"]: r["prob"] for r in rows}
    # 反证: 旧口径这里存的是 0.39 (P(是)), 而答案是「否」
    assert by == {"has_specific_place": 0.61, "product_role": 0.45, "own_experience": None}, by
    assert "0.61" in ing.rows_to_sql(rows), "SQL 那条路也要是新口径"
    print("  ✓ prob = 所选答案的概率: 答「否」存 1 − P(是) (0.39 → 0.61); 选择题第一名照存; 对不上 / 越界 → 不写")


def check_xlsx_roundtrip(bank: dict) -> None:
    import openpyxl
    B = _import_builder()
    qs = fb.llm_questions(bank)
    skipped = [q["id"] for q in qs][:11]
    rows = [{"note_id": "X_1", "block": 4, "title": "标题", "body": "正文", "skipped": skipped},
            {"note_id": "X_2", "block": 1, "title": None, "body": "正文二", "skipped": ["title_is_question"]}]
    man = {"X_1": {"annotators": "C", "skipped_questions": gl.join_skipped(skipped)},
           "X_2": {"annotators": "AC", "skipped_questions": gl.join_skipped(["title_is_question"])}}
    with tempfile.TemporaryDirectory() as d:
        path = Path(d) / "C.xlsx"
        B.build_workbook("C", rows, bank, path)
        wb = openpyxl.load_workbook(path)
        ws = wb["标注"]
        header = [c.value for c in ws[1]]
        qcol = {i: next(q for q in qs if str(h).split("\n", 1)[1] == gl.SHORT[q["id"]]) for i, h in enumerate(header)
                if h and "\n" in str(h) and not str(h).startswith("拿不准")}
        assert len(qcol) == 20, len(qcol)
        for r in range(2, ws.max_row + 1):
            for i, q in qcol.items():
                cell = ws.cell(row=r, column=i + 1)
                if cell.value != ing.SKIP_MARK:
                    cell.value = fb.closed_set(q)[0]
        wb.save(path)
        who, notes = ing.read_sheet(path, bank)
        assert who == "C" and {n["note_id"] for n in notes} == {"X_1", "X_2"}
        assert not ing.validate(path, who, notes, man, bank)
        got, blanks = ing.build_rows(notes, bank, "jev:t-C", "gate1-t")
        assert len(got) == 40 - 11 - 1 and not blanks, (len(got), blanks)
        # 灰格里填上答案再收 → 11 条错
        ws = wb["标注"]
        for r in range(2, ws.max_row + 1):
            for i, q in qcol.items():
                if ws.cell(row=r, column=i + 1).value == ing.SKIP_MARK and ws.cell(row=r, column=len(header)).value == "X_1":
                    ws.cell(row=r, column=i + 1).value = "否" if q["type"] == "bool" else fb.closed_set(q)[0]
        wb.save(path)
        who, notes = ing.read_sheet(path, bank)
        errs = ing.validate(path, who, notes, man, bank)
        assert len(errs) == 11, errs
    print("  ✓ xlsx 往返: build_workbook 灰 11 格 + 1 格 → 填白格 → read_sheet/validate/build_rows 28 行; 灰格填答案 → 11 条错")


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", action="store_true", help="再做一遍真 xlsx 往返 (要 openpyxl)")
    a = ap.parse_args(argv)
    bank = fb.load_bank()
    check_split_join()
    check_real_manifest(bank)
    check_prob(bank)
    if a.xlsx:
        check_xlsx_roundtrip(bank)
    # 名单级往返放最后: 没装 openpyxl 时它会往 sys.modules 塞哑壳, 别让 --xlsx 那一节拿到哑壳
    check_manifest_roundtrip(bank)
    print("✓ 闸一名单灰格往返 + prob 口径 (D-085): 全过" + ("" if a.xlsx else " (xlsx 往返用 --xlsx, 要 openpyxl)"))
    return 0


if __name__ == "__main__":
    sys.exit(main())
