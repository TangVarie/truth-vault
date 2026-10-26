#!/usr/bin/env python3
"""build_gate1_review_sheet.py —— 副题库闸一的 owner 裁决表 (D-128)。只在本地做, 文件直接交给 owner。

输入:
  · 格子表 (gate1_models.py --cells-out): kind=disagree 的格给 owner 裁, kind=spot 的 20 + 20 格放「抽查」页复核;
  · 样本名单 (gate1_sample.py 出的那份, 进了仓库): 每篇的 raw_md5;
  · 正文 JSON (--print-sql 拿去只读 MCP 跑, 存进 scratchpad): [{note_id, project_id, title, raw_content}, …];
    每篇 md5(raw_content) 必须等于名单里的 raw_md5, 对不上就不做 (正文变了 = 和模型看到的不是一篇, D-079 ⑤)。
输出: 一份 xlsx。⚠️ 它有正文, **不进仓库、不进 GitHub / Actions artifact / Artifact 页面**: 输出路径在仓库里且没被
  .gitignore 盖住就拒; 正文 JSON 同理。

    cd scripts && python build_gate1_review_sheet.py --cells ../data-analysis/gate1-reader-cells-<日期>.csv --print-sql
    python build_gate1_review_sheet.py --cells … --texts /path/to/scratchpad/texts.json --out /path/to/scratchpad/闸一裁决.xlsx

表的样子 (owner 2026-09-22:「尽量简单」):
  · 每题一页, 一行一格: 标题 / 正文 / 结尾那段 (代码切的, 同模型看到的) / 甲 乙 (丙) = 三趟给出的【不同】答案,
    哪个是甲按 sha256(种子:cell_id:答案) 排, 不写模型名、不写几趟选了它; 「你的判断」是下拉 (全部选项 + 拿不准),
    「理由」可不填 (不进库); 隐藏列 cell_id 对账用。没有公式。
  · 「抽查」页: 两题各 20 格三趟一致的, 显示那个一致的答案, 同样下拉复核 (推翻 > 2 个这题闸一不过, D-128)。
  · 「题目说明」页: 题面、看哪一段、每个选项的意思和例子 (就是 yaml 里给模型的那些)。
openpyxl 只在真做表时导入 (不在 requirements 里); 计划 / 校验 / 路径检查这些纯函数不装它也能跑 (CI 跑这些)。
"""
from __future__ import annotations

import argparse
import hashlib
import json
import subprocess
import sys
from pathlib import Path

import yaml

import feature_bank as fb
from gate1_agreement import load_pivot
from gate1_labels import SHORT
from gate1_models import Gate1Refused, load_secondary, read_cells
from gate1_sample import SampleRefused, read_manifest

ROOT = Path(__file__).resolve().parent.parent
SEED = "fqr-g1-2026-10-26:review"
UNSURE = "拿不准"
CANDIDATE_MARKS = ("甲", "乙", "丙")
SPOT_SHEET = "抽查"
HELP_SHEET = "怎么裁"
SPEC_SHEET = "题目说明"
TRUNC_NOTE = "\n\n……（后面截掉了，模型也只看到这里）"
SCOPE_ZH = {"title": "只看标题", "first_sentence": "只看正文第一句", "last_para": "只看正文最后一段（不算结尾的 #话题 和 @）",
            "body": "看正文", "full": "看标题 + 正文"}


class ReviewRefused(Exception):
    pass


def _md5(s: str) -> str:
    return hashlib.md5(s.encode("utf-8")).hexdigest()


def _rank(*parts: str) -> str:
    return hashlib.sha256(":".join(parts).encode("utf-8")).hexdigest()


def sheet_name(qid: str) -> str:
    return SHORT[qid]


def options(q: dict) -> list[str]:
    """下拉: 全部选项 + 拿不准 (Excel 列表 ≤ 255 字符)。"""
    return list(fb.closed_set(q)) + [UNSURE]


def check_private_path(path: Path, what: str) -> None:
    """有正文的文件: 在仓库里就必须被 .gitignore 盖住, 否则拒 (git check-ignore 说了算)。"""
    p = path.resolve()
    try:
        p.relative_to(ROOT)
    except ValueError:
        return
    try:
        r = subprocess.run(["git", "-C", str(ROOT), "check-ignore", "-q", "--", str(p)], capture_output=True, timeout=30)
    except (OSError, subprocess.SubprocessError) as exc:
        raise ReviewRefused(f"{what} {path} 在仓库里, 又查不了 .gitignore ({exc}) —— 放到 scratchpad") from exc
    if r.returncode != 0:
        raise ReviewRefused(f"{what} {path} 在仓库里且没被 .gitignore 盖住 —— 带正文的文件不进仓库, 放到 scratchpad")


def texts_sql(note_ids: list[str]) -> str:
    ids = ", ".join(f"'{n}'" for n in sorted(set(note_ids)))
    return ("-- build_gate1_review_sheet.py --print-sql · 只读 · 结果存进 scratchpad, 别存进仓库\n"
            "select json_agg(json_build_object('note_id', note_id, 'project_id', project_id, 'title', title, "
            "'raw_content', raw_content) order by note_id)::text as pivot\n"
            f"from truth_vault.notes where note_id in ({ids});")


def load_texts(path: Path, needed: set[str], manifest: list[dict]) -> dict[str, dict]:
    check_private_path(path, "正文 JSON")
    md5 = {m["note_id"]: m["raw_md5"] for m in manifest}
    got = {}
    for r in load_pivot(path):
        if isinstance(r, dict) and r.get("note_id") in needed:
            got[r["note_id"]] = r
    missing = sorted(needed - set(got))
    if missing:
        raise ReviewRefused(f"正文 JSON 缺 {len(missing)} 篇, 如 {missing[:3]}")
    bad = sorted(n for n, r in got.items() if _md5(r.get("raw_content") or "") != md5.get(n))
    if bad:
        raise ReviewRefused(f"{len(bad)} 篇正文 md5 和名单对不上 (如 {bad[:3]}) —— 正文改过了, 不是模型看到的那篇, 不做表")
    return got


def _title_mode(project_id: str, cache: dict) -> str:
    if project_id not in cache:
        m = yaml.safe_load((ROOT / "mappings" / f"{project_id}.yaml").read_text(encoding="utf-8")) or {}
        cache[project_id] = m.get("title_extraction") or "none"
    return cache[project_id]


def plan_review(cells: list[dict], bank: dict, texts: dict[str, dict], manifest: list[dict], seed: str = SEED) -> dict:
    """→ {"pages": {qid: [行]}, "spot": [行]}。行里只有 owner 要看的东西, 没有模型名 / 趟名。"""
    idx = fb.question_index(bank)
    mf = {m["note_id"]: m for m in manifest}
    modes: dict = {}
    pages: dict[str, list[dict]] = {q["id"]: [] for q in fb.llm_questions(bank)}
    spot: list[dict] = []
    for c in cells:
        if c["kind"] not in ("disagree", "spot"):
            continue
        if c["note_id"] not in mf:
            raise ReviewRefused(f"{c['cell_id']}: 不在名单里")
        if c["question_id"] not in idx:
            raise ReviewRefused(f"{c['cell_id']}: 不是本库的题")
        t = texts[c["note_id"]]
        sp = fb.build_spans(t.get("raw_content") or "", mode=_title_mode(mf[c["note_id"]]["project_id"], modes),
                            title_col=t.get("title"))
        answers = sorted({c[L] for L in ("A", "B", "H") if c.get(L) is not None},
                         key=lambda v: _rank(seed, c["cell_id"], v))
        row = {"cell_id": c["cell_id"], "question_id": c["question_id"], "title": sp["title"] or "（没有标题）",
               "body": (sp["body"] or "") + (TRUNC_NOTE if sp["_truncated"] else ""),
               "last_para": sp["last_para"] or "（没有）", "options": options(idx[c["question_id"]])}
        if c["kind"] == "disagree":
            if len(answers) > len(CANDIDATE_MARKS):
                raise ReviewRefused(f"{c['cell_id']}: {len(answers)} 个不同答案, 表只放得下 {len(CANDIDATE_MARKS)} 个")
            row["candidates"] = list(zip(CANDIDATE_MARKS, answers))
            pages[c["question_id"]].append(row)
        else:
            if len(answers) != 1:
                raise ReviewRefused(f"{c['cell_id']}: 抽查格应是三趟一致的一个答案, 实得 {answers}")
            row["model_answer"] = answers[0]
            spot.append(row)
    for q in pages:
        pages[q].sort(key=lambda r: _rank(seed, r["cell_id"]))
    spot.sort(key=lambda r: (list(pages).index(r["question_id"]), _rank(seed, r["cell_id"])))
    return {"pages": pages, "spot": spot}


def build_workbook(plan: dict, bank: dict, out: Path) -> None:
    from openpyxl import Workbook                                       # 只在真做表时要
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    check_private_path(out, "裁决表")
    f_body, f_bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    f_head = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    fill_head, fill_todo = PatternFill("solid", fgColor="44546A"), PatternFill("solid", fgColor="FFF2CC")
    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap, center = Alignment(wrap_text=True, vertical="top"), Alignment(horizontal="center", vertical="center", wrap_text=True)
    idx = fb.question_index(bank)
    n_dis = sum(len(v) for v in plan["pages"].values())

    wb = Workbook()
    ws0 = wb.active
    ws0.title = HELP_SHEET
    ws0.column_dimensions["B"].width = 100
    steps = [f"读者提问方式 · 闸一裁决 · 不一致 {n_dis} 格 + 抽查 {len(plan['spot'])} 格",
             "",
             "1. 每道题一页。每一行是一篇笔记的一格：几次判读给出的答案不一样，列成 甲 / 乙（/ 丙），顺序是随机的，和谁答的无关。",
             "2. 先读标题、正文和「结尾那段」，再看页顶的题目（「题目说明」页有每个选项的意思和例子），在黄色格子的下拉里选你认为对的那个。",
             "   可以选甲乙以外的选项。一时定不了先选「拿不准」，回头再改：每一格都要选一个答案，有一格是「拿不准」或空着，这道题就出不了结论。",
             "3. 「理由」可以不填，也不会存进库。",
             "4. 「抽查」页：几次判读答案一致的格子随机抽的，一题 20 个。同意就选同一个答案，不同意就选你认为对的。",
             "5. 只看表里给的东西判，不用查别的。改完把文件直接发回，不要上传到别处。"]
    for i, s in enumerate(steps, start=1):
        ws0.cell(row=i, column=2, value=s).font = f_bold if i == 1 else f_body
        ws0.cell(row=i, column=2).alignment = wrap

    def header(ws, heads, widths):
        for c, (h, w) in enumerate(zip(heads, widths), 1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font, cell.fill, cell.alignment, cell.border = f_head, fill_head, center, box
            ws.column_dimensions[_col(c)].width = w
        ws.column_dimensions[_col(len(heads))].hidden = True          # cell_id
        ws.freeze_panes = "B2"

    def text_cells(ws, ri, vals):
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=ri, column=c, value=v)
            cell.font, cell.alignment, cell.border = f_body, wrap, box

    for qid, rows in plan["pages"].items():
        ws = wb.create_sheet(sheet_name(qid))
        heads = ["序号", "标题", "正文", "结尾那段", "甲", "乙", "丙", "你的判断", "理由（可不填）", "cell_id"]
        header(ws, heads, [5, 18, 56, 28, 10, 10, 10, 12, 20, 4])
        dv = DataValidation(type="list", formula1='"' + ",".join(options(idx[qid])) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle="请从下拉里选", error="请从下拉选项里选一个")
        ws.add_data_validation(dv)
        for ri, r in enumerate(rows, start=2):
            cand = dict(r["candidates"])
            text_cells(ws, ri, [ri - 1, r["title"], r["body"], r["last_para"], cand.get("甲"), cand.get("乙"), cand.get("丙")])
            pick = ws.cell(row=ri, column=8)
            pick.fill, pick.border, pick.alignment = fill_todo, box, center
            dv.add(pick.coordinate)
            ws.cell(row=ri, column=9).border = box
            ws.cell(row=ri, column=10, value=r["cell_id"])
            ws.row_dimensions[ri].height = min(409, max(60, 14 * (len(r["body"]) // 28 + 2)))

    ws = wb.create_sheet(SPOT_SHEET)
    header(ws, ["序号", "题", "标题", "正文", "结尾那段", "判读的答案", "你的判断", "理由（可不填）", "cell_id"],
           [5, 12, 18, 56, 28, 12, 12, 20, 4])
    dvs = {}
    for qid in plan["pages"]:
        dvs[qid] = DataValidation(type="list", formula1='"' + ",".join(options(idx[qid])) + '"', allow_blank=True,
                                  showErrorMessage=True, errorTitle="请从下拉里选", error="请从下拉选项里选一个")
        ws.add_data_validation(dvs[qid])
    for ri, r in enumerate(plan["spot"], start=2):
        text_cells(ws, ri, [ri - 1, SHORT[r["question_id"]], r["title"], r["body"], r["last_para"], r["model_answer"]])
        pick = ws.cell(row=ri, column=7)
        pick.fill, pick.border, pick.alignment = fill_todo, box, center
        dvs[r["question_id"]].add(pick.coordinate)
        ws.cell(row=ri, column=8).border = box
        ws.cell(row=ri, column=9, value=r["cell_id"])
        ws.row_dimensions[ri].height = min(409, max(60, 14 * (len(r["body"]) // 28 + 2)))

    w3 = wb.create_sheet(SPEC_SHEET)
    for c, (h, wd) in enumerate(zip(["题", "完整问题", "看哪一段", "选项 / 算「是」", "算「否」"], (12, 40, 22, 70, 50)), 1):
        cell = w3.cell(row=1, column=c, value=h)
        cell.font, cell.fill, cell.alignment = f_head, fill_head, center
        w3.column_dimensions[_col(c)].width = wd
    for i, q in enumerate(fb.llm_questions(bank), start=2):
        if q["type"] == "choice":
            yes = "\n".join(f"· {o['value']}：{o.get('means', '')}（例：{o.get('example', '')}）" for o in q["options"])
            no = "（单选，选最贴近的一项）"
        else:
            yes = q.get("yes_if", "") + "\n例：" + "；".join(q.get("yes_examples") or [])
            no = q.get("no_if", "") + "\n例：" + "；".join(q.get("no_examples") or [])
        text_cells(w3, i, [SHORT[q["id"]], q["ask"], SCOPE_ZH.get(q["scope"], q["scope"]), yes, no])
    wb.active = 0
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def _col(i: int) -> str:
    s = ""
    while i:
        i, rem = divmod(i - 1, 26)
        s = chr(65 + rem) + s
    return s


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--cells", type=Path, required=True)
    ap.add_argument("--manifest", type=Path, default=ROOT / "data-analysis" / "gate1-reader-sample-2026-10-26.csv")
    ap.add_argument("--bank", default="reader")
    ap.add_argument("--texts", type=Path)
    ap.add_argument("--out", type=Path)
    ap.add_argument("--seed", default=SEED)
    ap.add_argument("--print-sql", action="store_true", help="打印取正文的只读 SQL (只取要裁 / 抽查的那几篇)")
    a = ap.parse_args(argv)
    try:
        bank = load_secondary(a.bank)
        cells = read_cells(a.cells)
        manifest = read_manifest(a.manifest)
        needed = {c["note_id"] for c in cells if c["kind"] in ("disagree", "spot")}
        if a.print_sql:
            print(texts_sql(sorted(needed)))
            return 0
        if not (a.texts and a.out):
            ap.error("要 --texts 和 --out (或 --print-sql)")
        check_private_path(a.out, "裁决表")
        plan = plan_review(cells, bank, load_texts(a.texts, needed, manifest), manifest, a.seed)
        build_workbook(plan, bank, a.out)
    except (ReviewRefused, Gate1Refused, SampleRefused) as exc:
        print(f"拒绝: {exc}", file=sys.stderr)
        return 2
    n = {q: len(v) for q, v in plan["pages"].items()}
    print(f"写到 {a.out}: 不一致 {n} · 抽查 {len(plan['spot'])} 格。直接交给 owner, 不进仓库 / GitHub。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
